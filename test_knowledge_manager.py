import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook

from knowledge_manager import KnowledgeManager


class KnowledgeManagerTests(unittest.TestCase):
    def setUp(self):
        self.temp_directory = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_directory.name)
        self.manager = KnowledgeManager(self.root / "assistant.db")

    def tearDown(self):
        self.temp_directory.cleanup()

    def test_add_list_search_and_deduplicate(self):
        source = self.root / "notes.md"
        source.write_text("项目目标是搭建可靠的个人 AI 助手。", encoding="utf-8")
        first = self.manager.add_document(source, "AI助手")
        second = self.manager.add_document(source, "AI助手")
        self.assertTrue(first["added"])
        self.assertFalse(second["added"])
        self.assertEqual(len(self.manager.list_documents()), 1)
        results = self.manager.search("可靠", "AI助手")
        self.assertEqual(results[0]["title"], "notes")

    def test_rejects_hidden_unsupported_and_empty_files(self):
        hidden = self.root / ".env"
        hidden.write_text("SECRET=x", encoding="utf-8")
        with self.assertRaises(ValueError):
            self.manager.add_document(hidden)
        unsupported = self.root / "notes.json"
        unsupported.write_text("{}", encoding="utf-8")
        with self.assertRaises(ValueError):
            self.manager.add_document(unsupported)
        empty = self.root / "empty.txt"
        empty.write_text("", encoding="utf-8")
        with self.assertRaises(ValueError):
            self.manager.add_document(empty)

    def test_imports_word_and_excel(self):
        word_path = self.root / "notes.docx"
        word = Document()
        word.add_paragraph("Word 文档里的项目结论")
        word.save(word_path)
        excel_path = self.root / "plan.xlsx"
        workbook = Workbook()
        workbook.active.append(["事项", "负责人"])
        workbook.active.append(["完成联网功能", "小林"])
        workbook.save(excel_path)
        workbook.close()

        self.assertTrue(self.manager.add_document(word_path, "办公资料")["added"])
        self.assertTrue(self.manager.add_document(excel_path, "办公资料")["added"])
        self.assertIn("项目结论", self.manager.search("项目结论")[0]["content"])
        self.assertIn("完成联网功能", self.manager.search("完成联网功能")[0]["content"])

    def test_full_text_search_prioritises_more_relevant_chunk(self):
        weak = self.root / "weak.txt"
        strong = self.root / "strong.txt"
        weak.write_text("苹果手机偶尔被提到。这里还有很多无关内容。", encoding="utf-8")
        strong.write_text("苹果手机选购：苹果手机型号、苹果手机价格。", encoding="utf-8")
        self.manager.add_document(weak)
        self.manager.add_document(strong)
        results = self.manager.search("苹果手机")
        self.assertEqual(results[0]["title"], "strong")

    def test_rejects_office_archive_with_excessive_expanded_size(self):
        source = Path(self.temp_directory.name) / "oversized.docx"
        with zipfile.ZipFile(source, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("word/document.xml", "x" * 1024)
        with patch("knowledge_manager.MAX_EXPANDED_DOCUMENT_BYTES", 100):
            with self.assertRaisesRegex(ValueError, "解压后"):
                self.manager.add_document(source)

    def test_project_filter_separates_context(self):
        one = self.root / "one.txt"
        two = self.root / "two.txt"
        one.write_text("共同关键词，甲项目资料", encoding="utf-8")
        two.write_text("共同关键词，乙项目资料", encoding="utf-8")
        self.manager.add_document(one, "甲")
        self.manager.add_document(two, "乙")
        results = self.manager.search("共同关键词", project="乙")
        self.assertEqual([item["project"] for item in results], ["乙"])

    def test_same_content_can_exist_in_two_projects_and_be_deleted(self):
        source = self.root / "shared.txt"
        source.write_text("相同资料内容", encoding="utf-8")
        first = self.manager.add_document(source, "甲")
        second = self.manager.add_document(source, "乙")
        self.assertNotEqual(first["document_id"], second["document_id"])
        self.assertTrue(self.manager.delete_document(first["document_id"]))
        self.assertEqual(self.manager.search("相同资料内容", "甲"), [])
        self.assertEqual(len(self.manager.search("相同资料内容", "乙")), 1)

    def test_percent_and_underscore_are_literal_search_terms(self):
        source = self.root / "symbols.txt"
        source.write_text("普通内容，不包含通配符", encoding="utf-8")
        self.manager.add_document(source)
        self.assertEqual(self.manager.search("%"), [])
        self.assertEqual(self.manager.search("_"), [])


if __name__ == "__main__":
    unittest.main()
