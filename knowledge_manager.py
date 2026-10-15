"""本地资料库：显式导入常用文档，按项目保存并检索。"""

from __future__ import annotations

import hashlib
import sqlite3
import zipfile
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path

from conversation_manager import DEFAULT_DATABASE

ALLOWED_SUFFIXES = {".csv", ".docx", ".md", ".pdf", ".txt", ".xlsx"}
MAX_DOCUMENT_BYTES = 20_000_000
MAX_EXTRACTED_CHARS = 2_000_000
MAX_EXPANDED_DOCUMENT_BYTES = 100_000_000
MAX_ARCHIVE_MEMBERS = 10_000
CHUNK_SIZE = 1600
CHUNK_OVERLAP = 200


def _utc_now():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


class KnowledgeManager:
    """在本地 SQLite 中保存用户明确导入的 UTF-8 文本资料。"""

    def __init__(self, database_path=DEFAULT_DATABASE):
        self.database_path = Path(database_path)
        self.database_path.parent.mkdir(parents=True, exist_ok=True)
        self._initialise()

    @contextmanager
    def _connect(self):
        connection = sqlite3.connect(self.database_path)
        connection.row_factory = sqlite3.Row
        connection.execute("PRAGMA foreign_keys = ON")
        try:
            yield connection
            connection.commit()
        except Exception:
            connection.rollback()
            raise
        finally:
            connection.close()

    def _initialise(self):
        with self._connect() as connection:
            connection.executescript(
                """
                CREATE TABLE IF NOT EXISTS knowledge_documents (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    source_path TEXT NOT NULL,
                    title TEXT NOT NULL,
                    project TEXT NOT NULL,
                    content_hash TEXT NOT NULL,
                    created_at TEXT NOT NULL
                    ,UNIQUE(content_hash, project)
                );

                CREATE TABLE IF NOT EXISTS knowledge_chunks (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    document_id INTEGER NOT NULL,
                    position INTEGER NOT NULL,
                    content TEXT NOT NULL,
                    FOREIGN KEY (document_id) REFERENCES knowledge_documents(id)
                        ON DELETE CASCADE
                );

                CREATE INDEX IF NOT EXISTS knowledge_chunks_document
                ON knowledge_chunks(document_id, position);

                CREATE VIRTUAL TABLE IF NOT EXISTS knowledge_fts USING fts5(
                    content,
                    document_id UNINDEXED,
                    position UNINDEXED,
                    tokenize='trigram'
                );

                CREATE TRIGGER IF NOT EXISTS knowledge_chunks_fts_insert
                AFTER INSERT ON knowledge_chunks BEGIN
                    INSERT INTO knowledge_fts(rowid, content, document_id, position)
                    VALUES (new.id, new.content, new.document_id, new.position);
                END;

                CREATE TRIGGER IF NOT EXISTS knowledge_chunks_fts_delete
                AFTER DELETE ON knowledge_chunks BEGIN
                    DELETE FROM knowledge_fts WHERE rowid = old.id;
                END;

                CREATE TRIGGER IF NOT EXISTS knowledge_chunks_fts_update
                AFTER UPDATE ON knowledge_chunks BEGIN
                    DELETE FROM knowledge_fts WHERE rowid = old.id;
                    INSERT INTO knowledge_fts(rowid, content, document_id, position)
                    VALUES (new.id, new.content, new.document_id, new.position);
                END;
                """
            )
            connection.execute(
                """
                INSERT INTO knowledge_fts(rowid, content, document_id, position)
                SELECT c.id, c.content, c.document_id, c.position
                FROM knowledge_chunks c
                LEFT JOIN knowledge_fts f ON f.rowid = c.id
                WHERE f.rowid IS NULL
                """
            )

    @staticmethod
    def _validate_source(path):
        source = Path(path).expanduser().resolve()
        if not source.is_file():
            raise ValueError("找不到这个资料文件。")
        if source.name.startswith(".") or source.suffix.lower() not in ALLOWED_SUFFIXES:
            raise ValueError("资料库接受非隐藏的 PDF、Word、Excel、CSV、Markdown 和 TXT 文件。")
        if source.stat().st_size > MAX_DOCUMENT_BYTES:
            raise ValueError("单份资料不能超过 20 MB。")
        return source

    @staticmethod
    def _validate_office_archive(source):
        try:
            with zipfile.ZipFile(source) as archive:
                members = archive.infolist()
                expanded = sum(member.file_size for member in members)
        except zipfile.BadZipFile as error:
            raise ValueError("Office 资料压缩结构损坏。") from error
        if len(members) > MAX_ARCHIVE_MEMBERS:
            raise ValueError("Office 资料内部文件数量异常，已拒绝导入。")
        if expanded > MAX_EXPANDED_DOCUMENT_BYTES:
            raise ValueError("Office 资料解压后超过 100 MB，已拒绝导入。")

    @staticmethod
    def _extract_text(source):
        suffix = source.suffix.lower()
        if suffix in {".md", ".txt", ".csv"}:
            try:
                content = source.read_text(encoding="utf-8")
            except UnicodeDecodeError as error:
                raise ValueError("文本资料必须是 UTF-8 编码。") from error
        elif suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(source)
            if reader.is_encrypted:
                raise ValueError("暂不支持加密 PDF。")
            content = "\n\n".join(page.extract_text() or "" for page in reader.pages)
        elif suffix == ".docx":
            KnowledgeManager._validate_office_archive(source)
            from docx import Document

            document = Document(source)
            lines = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    lines.append("\t".join(cell.text for cell in row.cells))
            content = "\n".join(lines)
        elif suffix == ".xlsx":
            KnowledgeManager._validate_office_archive(source)
            from openpyxl import load_workbook

            workbook = load_workbook(source, read_only=True, data_only=True)
            lines = []
            cell_count = 0
            try:
                for sheet in workbook.worksheets:
                    lines.append(f"## 工作表：{sheet.title}")
                    for row in sheet.iter_rows(values_only=True):
                        values = ["" if value is None else str(value) for value in row]
                        if any(values):
                            lines.append("\t".join(values))
                            cell_count += sum(bool(value) for value in values)
                            if cell_count > 100_000:
                                raise ValueError("Excel 非空单元格超过 10 万个，请拆分后导入。")
            finally:
                workbook.close()
            content = "\n".join(lines)
        else:
            raise ValueError("不支持这个资料格式。")
        if len(content) > MAX_EXTRACTED_CHARS:
            raise ValueError("资料提取后的文字超过 200 万字符，请拆分后导入。")
        return content

    @staticmethod
    def _chunks(text):
        text = text.strip()
        if not text:
            raise ValueError("资料文件是空的。")
        step = CHUNK_SIZE - CHUNK_OVERLAP
        return [text[index : index + CHUNK_SIZE] for index in range(0, len(text), step)]

    def add_document(self, path, project="default"):
        source = self._validate_source(path)
        project = str(project).strip() or "default"
        if len(project) > 80:
            raise ValueError("项目名称不能超过 80 个字符。")
        content = self._extract_text(source)
        chunks = self._chunks(content)
        digest = hashlib.sha256(content.encode("utf-8")).hexdigest()

        with self._connect() as connection:
            existing = connection.execute(
                "SELECT id FROM knowledge_documents WHERE content_hash = ? AND project = ?",
                (digest, project),
            ).fetchone()
            if existing:
                return {"document_id": existing["id"], "added": False}
            cursor = connection.execute(
                """
                INSERT INTO knowledge_documents
                    (source_path, title, project, content_hash, created_at)
                VALUES (?, ?, ?, ?, ?)
                """,
                (str(source), source.stem, project, digest, _utc_now()),
            )
            document_id = cursor.lastrowid
            connection.executemany(
                """
                INSERT INTO knowledge_chunks (document_id, position, content)
                VALUES (?, ?, ?)
                """,
                [(document_id, position, chunk) for position, chunk in enumerate(chunks)],
            )
        return {"document_id": document_id, "added": True}

    def list_documents(self):
        with self._connect() as connection:
            rows = connection.execute(
                """
                SELECT id, title, project, source_path, created_at
                FROM knowledge_documents ORDER BY id DESC
                """
            ).fetchall()
        return [dict(row) for row in rows]

    def delete_document(self, document_id):
        try:
            document_id = int(document_id)
        except (TypeError, ValueError) as error:
            raise ValueError("资料编号必须是数字。") from error
        if document_id <= 0:
            raise ValueError("资料编号必须是正整数。")
        with self._connect() as connection:
            cursor = connection.execute(
                "DELETE FROM knowledge_documents WHERE id = ?", (document_id,)
            )
            return cursor.rowcount == 1

    def search(self, query, project="", limit=5):
        query = str(query).strip()
        project = str(project).strip()
        if not query:
            raise ValueError("检索词不能为空。")
        if not isinstance(limit, int) or not 1 <= limit <= 10:
            raise ValueError("检索数量必须在 1 到 10 之间。")
        if len(query) >= 3:
            escaped_fts = query.replace('"', '""')
            fts_parameters = [f'"{escaped_fts}"']
            fts_project_sql = ""
            if project:
                fts_project_sql = " AND d.project = ?"
                fts_parameters.append(project)
            fts_parameters.append(limit)
            with self._connect() as connection:
                rows = connection.execute(
                    f"""
                    SELECT d.id AS document_id, d.title, d.project,
                           c.position, c.content
                    FROM knowledge_fts
                    JOIN knowledge_chunks c ON c.id = knowledge_fts.rowid
                    JOIN knowledge_documents d ON d.id = c.document_id
                    WHERE knowledge_fts MATCH ?{fts_project_sql}
                    ORDER BY bm25(knowledge_fts), d.id DESC, c.position ASC
                    LIMIT ?
                    """,
                    fts_parameters,
                ).fetchall()
            if rows:
                return [dict(row) for row in rows]

        escaped_query = (
            query.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
        )
        like_query = f"%{escaped_query}%"
        parameters = [like_query]
        project_sql = ""
        if project:
            project_sql = " AND d.project = ?"
            parameters.append(project)
        parameters.append(limit)
        with self._connect() as connection:
            rows = connection.execute(
                f"""
                SELECT d.id AS document_id, d.title, d.project, c.position, c.content
                FROM knowledge_chunks c
                JOIN knowledge_documents d ON d.id = c.document_id
                WHERE c.content LIKE ? ESCAPE '\\'{project_sql}
                ORDER BY d.id DESC, c.position ASC
                LIMIT ?
                """,
                parameters,
            ).fetchall()
        return [dict(row) for row in rows]
